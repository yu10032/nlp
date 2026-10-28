import asyncio
import json
from io import BytesIO

import httpx
from docx import Document
from docx.oxml import OxmlElement
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from sqlalchemy import select
from sqlalchemy.orm import Session
from ta_backend.config import Settings
from ta_backend.main import create_app
from ta_backend.models import IngestionJob, Material, QuestionPrediction
from ta_backend.services.generation import ExternalAPIProvider, GeneratedAnswer

ROOT = "/api/v1/courses/demo-nlp"


def upload(
    client,
    filename="lecture.md",
    text="# 注意力机制\n\n注意力机制对值向量加权求和。",
    replaces=None,
):
    path = f"{ROOT}/documents" + (f"?replaces_id={replaces}" if replaces else "")
    response = client.post(path, files={"file": (filename, text.encode("utf-8"), "text/plain")})
    assert response.status_code == 202, response.text
    job = response.json()
    return job, client.get(f"{ROOT}/jobs/{job['job_id']}").json()


def make_question(client):
    response = client.post(
        f"{ROOT}/questions/match", json={"question_text": "注意力机制怎样加权求和？"}
    )
    assert response.status_code == 201
    return response.json()


def make_draft(client, question_id, citation_id="demo-chunk-2"):
    response = client.post(
        f"{ROOT}/drafts",
        json={
            "question_id": question_id,
            "title": "注意力机制答疑",
            "answer": "按匹配结果对值向量加权求和。",
            "node_ids": ["demo-attention"],
            "citation_ids": [citation_id],
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def test_docx_headings_tables_and_partial_formula_report(api):
    client, _ = api
    source = Document()
    source.add_heading("注意力机制", level=1)
    source.add_paragraph("根据查询与键的匹配结果，对值向量加权求和。")
    table = source.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "概念"
    table.cell(0, 1).text = "含义"
    table.cell(1, 0).text = "Query"
    table.cell(1, 1).text = "查询"
    source.add_paragraph()._p.append(OxmlElement("m:oMath"))
    output = BytesIO()
    source.save(output)
    response = client.post(f"{ROOT}/documents", files={"file": ("lecture.docx", output.getvalue())})
    assert response.status_code == 202
    material_id = response.json()["material_id"]
    docs = client.get(f"{ROOT}/documents").json()
    material = next(doc for doc in docs if doc["id"] == material_id)
    assert material["status"] == "partial"
    assert "公式" in material["warnings"][0]
    chunks = client.get(f"{ROOT}/documents/{material_id}/chunks").json()
    assert any(chunk["section_path"] == ["注意力机制"] for chunk in chunks)
    assert any("Query | 查询" in chunk["text"] for chunk in chunks)
    assert all(chunk["page_from"] is None for chunk in chunks)
    assert client.get(f"{ROOT}/documents/{material_id}/file").content == output.getvalue()


def test_disguised_files_and_scanned_pdf_do_not_create_empty_success(api):
    client, _ = api
    assert (
        client.post(f"{ROOT}/documents", files={"file": ("fake.docx", b"not a docx")}).status_code
        == 415
    )
    assert (
        client.post(f"{ROOT}/documents", files={"file": ("fake.txt", b"%PDF-1.0")}).status_code
        == 415
    )
    writer = PdfWriter()
    writer.add_blank_page(width=500, height=700)
    output = BytesIO()
    writer.write(output)
    response = client.post(f"{ROOT}/documents", files={"file": ("scan.pdf", output.getvalue())})
    assert response.status_code == 202
    job = client.get(f"{ROOT}/jobs/{response.json()['job_id']}").json()
    assert job["status"] == "failed"
    assert "OCR" in job["error"]


def test_material_update_invalidates_review_preserves_original_and_revisions(api):
    client, _ = api
    first, job = upload(client)
    assert job["status"] == "done"
    chunk = client.get(f"{ROOT}/documents/{first['material_id']}/chunks").json()[0]
    question = make_question(client)
    draft = make_draft(client, question["id"], chunk["id"])
    reviewed = client.post(
        f"{ROOT}/drafts/{draft['id']}/review",
        json={
            "expected_version": 1,
            "result": "approved",
            "reviewer": "测试教师",
        },
    )
    assert reviewed.status_code == 200
    assert client.get(f"{ROOT}/drafts/{draft['id']}/export").status_code == 200
    second, job = upload(
        client, text="# 注意力机制\n\n新版课程资料补充了缩放说明。", replaces=first["material_id"]
    )
    assert job["status"] == "done"
    saved = client.get(f"{ROOT}/drafts/{draft['id']}").json()
    assert saved["status"] == "stale_review"
    assert saved["reviews"][0]["revision"] == 1
    assert saved["revisions"][0]["snapshot"]["answer"] == draft["answer"]
    assert saved["citations"][0]["active"] is False
    assert client.get(f"{ROOT}/drafts/{draft['id']}/export").status_code == 409
    docs = client.get(f"{ROOT}/documents").json()
    assert next(doc for doc in docs if doc["id"] == second["material_id"])["version"] == 2
    assert (
        client.get(f"{ROOT}/documents/{first['material_id']}/chunks").json()[0]["text"]
        == chunk["text"]
    )


def test_edits_expire_review_and_conflicting_revision_is_rejected(api):
    client, _ = api
    draft = make_draft(client, make_question(client)["id"])
    client.post(
        f"{ROOT}/drafts/{draft['id']}/review",
        json={"expected_version": 1, "result": "approved", "reviewer": "教师"},
    )
    update = {
        "expected_version": 1,
        "title": draft["title"],
        "answer": "修订后的解释",
        "node_ids": draft["node_ids"],
        "citation_ids": draft["citation_ids"],
    }
    response = client.patch(f"{ROOT}/drafts/{draft['id']}", json=update)
    assert response.json()["version"] == 2
    assert response.json()["status"] == "stale_review"
    assert {item["version"] for item in response.json()["revisions"]} == {1, 2}
    assert client.patch(f"{ROOT}/drafts/{draft['id']}", json=update).status_code == 409


def test_predictions_survive_human_correction_and_empty_confirmation(api):
    client, engine = api
    question = make_question(client)
    assert question["predictions"]
    changed = client.post(f"{ROOT}/questions/{question['id']}/links", json={"node_ids": []})
    assert changed.json()["needs_review"] is False
    with Session(engine) as session:
        assert (
            session.scalar(
                select(QuestionPrediction).where(QuestionPrediction.question_id == question["id"])
            )
            is not None
        )
    unknown = client.post(
        f"{ROOT}/questions/match", json={"question_text": "考试几点开始？"}
    ).json()
    assert unknown["predictions"] == []
    evidence = client.post(
        f"{ROOT}/evidence/retrieve", json={"question_text": "考试几点开始？"}
    ).json()
    assert evidence["missing_evidence"] is True


def test_course_boundaries_reject_chunks_nodes_and_drafts(api):
    client, _ = api
    other = client.post("/api/v1/courses", json={"title": "另一门课程"}).json()["id"]
    path = f"/api/v1/courses/{other}"
    assert client.get(f"{path}/documents/demo-material/chunks").status_code == 404
    response = client.post(
        f"{path}/nodes", json={"name": "跨课程节点", "evidence_chunk_ids": ["demo-chunk-2"]}
    )
    assert response.status_code == 422
    assert (
        client.post(
            f"{path}/evidence/retrieve",
            json={"question_text": "注意力", "node_ids": ["demo-attention"]},
        ).status_code
        == 404
    )
    draft = make_draft(client, make_question(client)["id"])
    assert client.get(f"{path}/drafts/{draft['id']}").status_code == 404


def test_candidate_edges_require_valid_sources_and_prerequisites_do_not_cycle(api):
    client, _ = api
    payload = {
        "source_id": "demo-attention",
        "target_id": "demo-transformer",
        "relation": "prerequisite_of",
        "chunk_id": "demo-chunk-2",
    }
    edge = client.post(f"{ROOT}/edges", json=payload)
    assert edge.status_code == 201
    assert edge.json()["status"] == "candidate"
    assert (
        client.post(
            f"{ROOT}/edges/{edge.json()['id']}/review", json={"result": "confirmed"}
        ).status_code
        == 200
    )
    reverse = client.post(
        f"{ROOT}/edges",
        json={**payload, "source_id": "demo-transformer", "target_id": "demo-attention"},
    ).json()
    assert (
        client.post(
            f"{ROOT}/edges/{reverse['id']}/review", json={"result": "confirmed"}
        ).status_code
        == 422
    )
    assert client.post(f"{ROOT}/edges", json={**payload, "chunk_id": "missing"}).status_code == 422


def test_disabled_and_invalid_generation_never_save_fake_drafts(api):
    client, _ = api
    question = make_question(client)
    payload = {
        "question_id": question["id"],
        "title": "答疑",
        "generate": True,
        "node_ids": ["demo-attention"],
    }
    response = client.post(f"{ROOT}/drafts", json=payload)
    assert response.status_code == 503
    assert response.json()["error"]["code"] == "generation_disabled"
    assert client.get(f"{ROOT}/drafts").json() == []


def test_external_adapter_sends_structured_request_without_real_network(monkeypatch):
    from ta_backend.services import generation

    original = httpx.AsyncClient
    captured = {}

    def handle(request):
        captured.update(json.loads(request.content))
        return httpx.Response(
            200,
            json={
                "choices": [
                    {
                        "message": {
                            "content": json.dumps(
                                {
                                    "answer": "测试响应",
                                    "citation_ids": ["chunk-1"],
                                    "missing_info": [],
                                }
                            )
                        }
                    }
                ],
                "usage": {"total_tokens": 10},
            },
        )

    monkeypatch.setattr(
        generation.httpx,
        "AsyncClient",
        lambda **kwargs: original(transport=httpx.MockTransport(handle), **kwargs),
    )
    settings = Settings(
        database_url="sqlite:///:memory:",
        llm_base_url="https://test.example/v1",
        llm_api_key="test-only",
        llm_model="test-model",
    )
    result, info = asyncio.run(
        ExternalAPIProvider(settings).generate_structured(
            "测试问题", [{"id": "chunk-1", "text": "测试原文"}]
        )
    )
    assert isinstance(result, GeneratedAnswer)
    assert captured["response_format"] == {"type": "json_object"}
    assert "tools" not in captured
    assert info["model_id"] == "test-model"


def test_restart_retains_documents_and_marks_interrupted_jobs(api, tmp_path):
    client, engine = api
    record, job = upload(client)
    assert job["status"] == "done"
    with Session(engine) as session:
        material = Material(
            course_id="demo-nlp", title="中断任务", status="ingesting", active=False
        )
        session.add(material)
        session.flush()
        pending = IngestionJob(course_id="demo-nlp", material_id=material.id, status="running")
        session.add(pending)
        session.commit()
        pending_id = pending.id
    settings = Settings(database_url=str(engine.url), storage_dir=tmp_path / "uploads")
    with TestClient(create_app(settings)) as restarted:
        assert restarted.get(f"{ROOT}/documents/{record['material_id']}/chunks").json()
        assert restarted.get(f"{ROOT}/jobs/{pending_id}").json()["status"] == "failed"
