import base64
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.oxml import parse_xml
from docx.shared import Cm
from ta_backend.services.lesson_export import append_prose
from test_teaching import ROOT, create, full_lesson_payload, quiz_payload, review

PIXEL = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aD1sAAAAASUVORK5CYII="
)


def test_word_export_does_not_drop_text_with_an_unclosed_math_delimiter():
    document = Document()
    append_prose(document, "$5 是教学假设中的成本，不是一个完整公式。")
    assert document.paragraphs[0].text == "$5 是教学假设中的成本，不是一个完整公式。"


def test_word_formula_and_embedded_image_have_readable_sources_and_scoped_preview(api):
    client, _ = api
    document = Document()
    document.add_heading("加权平均", 1)
    paragraph = document.add_paragraph("演示公式：")
    paragraph._p.append(
        parse_xml("""<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
        <m:f><m:num><m:r><m:t>a</m:t></m:r></m:num><m:den><m:r><m:t>b</m:t></m:r></m:den></m:f>
        </m:oMath>""")
    )
    document.add_picture(BytesIO(PIXEL), width=Cm(1))
    output = BytesIO()
    document.save(output)
    response = client.post(f"{ROOT}/documents", files={"file": ("formula.docx", output.getvalue())})
    assert response.status_code == 202
    material = response.json()["material_id"]
    chunks = client.get(f"{ROOT}/documents/{material}/chunks").json()
    assert any(r"$\frac{a}{b}$" in chunk["text"] for chunk in chunks)
    image_chunk = next(c for c in chunks if "asset:" in c["text"])
    relation = image_chunk["text"].split("asset:")[1].split(")")[0]
    image = client.get(f"{ROOT}/documents/{material}/images/{relation}")
    assert image.status_code == 200 and image.content == PIXEL
    assert image.headers["content-type"] == "image/png"
    assert client.get(f"{ROOT}/documents/{material}/images/rId999").status_code == 404
    other = client.post("/api/v1/courses", json={"title": "其他课程"}).json()["id"]
    assert (
        client.get(f"/api/v1/courses/{other}/documents/{material}/images/{relation}").status_code
        == 404
    )


def test_word_export_preserves_content_editable_equations_and_review_boundary(api):
    client, _ = api
    payload = full_lesson_payload()
    payload["body"]["segments"][1]["content"] += (
        "\n\n演算：$\\frac{a}{b} + x^{2}$。\n\n$$\\sum_{i=1}^{n} w_i v_i$$"
    )
    payload["body"]["segments"][1]["content"] += (
        "\n\n| 权重 | 输出 |\n| --- | --- |\n| 0.25 / 0.75 | $5$ |"
    )
    lesson = create(client, "lesson", payload)
    path = f"{ROOT}/teaching/lesson/{lesson['id']}"
    assert client.get(f"{path}/word").status_code == 409
    assert review(client, "lesson", lesson).status_code == 200
    output = client.get(f"{path}/word")
    assert output.status_code == 200, output.text if output.status_code != 200 else ""
    document = Document(BytesIO(output.content))
    text = "\n".join(p.text for p in document.paragraphs)
    for expected in [
        "教学目标",
        "课堂安排",
        "教师引导",
        "学生任务",
        "参考答案",
        "课后巩固",
        "课程依据",
    ]:
        assert expected in text
    with ZipFile(BytesIO(output.content)) as package:
        xml = package.read("word/document.xml").decode()
    assert "<m:f>" in xml and "<m:sSup>" in xml and "<m:nary>" in xml
    assert document.tables[0].cell(0, 0).text == "权重"
    assert document.tables[0].cell(1, 0).text == "0.25 / 0.75"
    assert client.get(f"{path}/print?expected_version={lesson['version']}").status_code == 200
    assert client.get(f"{path}/print?expected_version=999").status_code == 409
    changed = client.patch(path, json={**payload, "expected_version": lesson["version"]})
    assert changed.status_code == 200
    assert client.get(f"{path}/word").status_code == 409
    assert client.get(f"{path}/print?expected_version=2").status_code == 409


def test_student_export_uses_only_student_tasks_and_quiz_questions(api):
    client, _ = api
    quiz_data = quiz_payload()
    quiz_data["body"]["items"][0]["explanation"] = "只供教师查看的解析标记"
    quiz = create(client, "quiz", quiz_data)
    assert review(client, "quiz", quiz).status_code == 200
    payload = full_lesson_payload()
    payload["body"]["homework"] = "教师版混合参考答案，不应出现在学生版"
    payload["body"]["student_homework"] = "学生单独计算一组新权重的输出。"
    payload["body"]["segments"][0]["expected_response"] = "只供教师查看的参考回答"
    payload["body"]["segments"][1]["quiz_ids"] = [quiz["id"]]
    lesson = create(client, "lesson", payload)
    assert review(client, "lesson", lesson).status_code == 200
    output = client.get(f"{ROOT}/teaching/lesson/{lesson['id']}/word?audience=student")
    assert output.status_code == 200
    text = "\n".join(p.text for p in Document(BytesIO(output.content)).paragraphs)
    for hidden in [
        "只供教师查看",
        "教师版混合参考答案",
        "预期回答与参考答案",
        "教师引导与课堂活动",
    ]:
        assert hidden not in text
    for expected in ["学生版课堂讲义", "学生单独计算", "注意力机制如何汇总值向量", "加权求和"]:
        assert expected in text
    teacher = client.get(f"{ROOT}/teaching/lesson/{lesson['id']}/word?audience=teacher")
    teacher_text = "\n".join(p.text for p in Document(BytesIO(teacher.content)).paragraphs)
    assert "只供教师查看的解析标记" in teacher_text
