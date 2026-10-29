"""将已审核讲义排为 Word；公式保留为可编辑的 Office Math。"""

import re
from io import BytesIO
from xml.sax import SAXException

import latex2mathml.exceptions as math_errors
import mathml2omml
from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from latex2mathml.converter import convert

from ta_backend.services.teaching import LESSON_STAGES

MATH_ERRORS = (
    math_errors.NoAvailableTokensError,
    math_errors.NumeratorNotFoundError,
    math_errors.DenominatorNotFoundError,
    math_errors.ExtraLeftOrMissingRightError,
    math_errors.MissingSuperScriptOrSubscriptError,
    math_errors.DoubleSubscriptsError,
    math_errors.DoubleSuperscriptsError,
    math_errors.InvalidStyleForGenfracError,
    math_errors.MissingEndError,
    math_errors.InvalidAlignmentError,
    math_errors.InvalidWidthError,
    math_errors.LimitsMustFollowMathOperatorError,
    NotImplementedError,
    RuntimeError,
    SAXException,
)
MATH_PATTERN = r"(\$\$[\s\S]*?\$\$|\$[^$\n]+\$|\\\[[\s\S]*?\\\]|\\\([\s\S]*?\\\))"


class LessonExportError(ValueError):
    pass


def append_inline(paragraph, text: str):
    for piece in re.split(MATH_PATTERN, text):
        if not piece:
            continue
        if re.fullmatch(MATH_PATTERN, piece):
            width = 2 if piece.startswith(("$$", r"\(", r"\[")) else 1
            formula = piece[width:-width]
            try:
                xml = mathml2omml.convert(convert(formula))
            except MATH_ERRORS as error:
                raise LessonExportError(
                    f"此公式暂不能导出为 Word 公式，请修改后重试：{formula[:100]}"
                ) from error
            root = parse_xml(
                f'<root xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">{xml}</root>'
            )
            paragraph._p.append(root[0])
        else:
            for fragment in re.split(r"(\*\*.+?\*\*|`[^`]+`)", piece):
                if fragment.startswith("**") and fragment.endswith("**"):
                    paragraph.add_run(fragment[2:-2]).bold = True
                elif fragment.startswith("`") and fragment.endswith("`"):
                    paragraph.add_run(fragment[1:-1]).font.name = "Consolas"
                else:
                    paragraph.add_run(fragment)


def append_prose(document, text: str):
    # 公式块不按换行拆开，避免多行推导在导出时损坏。
    blocks = re.split(r"(\$\$[\s\S]*?\$\$|\\\[[\s\S]*?\\\])", text)
    for block in blocks:
        if re.fullmatch(r"\$\$[\s\S]*?\$\$|\\\[[\s\S]*?\\\]", block):
            append_inline(document.add_paragraph(), block)
            continue
        code = False
        lines = block.splitlines()
        index = 0
        while index < len(lines):
            line = lines[index]
            index += 1
            if line.lstrip().startswith("```"):
                code = not code
                continue
            if not line.strip():
                continue
            if code:
                paragraph = document.add_paragraph()
                paragraph.add_run(line).font.name = "Consolas"
                paragraph.paragraph_format.space_after = Pt(0)
                continue
            if index < len(lines) and re.fullmatch(
                r"\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?\s*", lines[index]
            ):

                def cells(value):
                    return [
                        cell.strip().replace(r"\|", "|")
                        for cell in re.split(r"(?<!\\)\|", value.strip().strip("|"))
                    ]

                headers = cells(line)
                table = document.add_table(rows=1, cols=len(headers))
                table.style = "Table Grid"
                repeat = OxmlElement("w:tblHeader")
                table.rows[0]._tr.get_or_add_trPr().append(repeat)
                for cell, value in zip(table.rows[0].cells, headers):
                    append_inline(cell.paragraphs[0], value)
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
                index += 1
                while index < len(lines) and "|" in lines[index] and lines[index].strip():
                    values = cells(lines[index])
                    if len(values) != len(headers):
                        raise LessonExportError("讲义中的 Markdown 表格列数不一致，请修正后导出")
                    for cell, value in zip(table.add_row().cells, values):
                        append_inline(cell.paragraphs[0], value)
                    index += 1
                continue
            heading = re.match(r"^(#{1,6})\s+(.+)$", line)
            if heading:
                append_inline(
                    document.add_heading(level=min(len(heading.group(1)) + 1, 4)), heading.group(2)
                )
            elif re.match(r"^\s*[-*]\s+", line):
                append_inline(
                    document.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*]\s+", "", line)
                )
            elif re.match(r"^\s*\d+[.)]\s+", line):
                append_inline(
                    document.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", line)
                )
            else:
                append_inline(document.add_paragraph(), line)


def export_lesson_word(
    content,
    citations: list[dict],
    node_names: dict[str, str],
    material_paths: dict[str, str],
    *,
    audience: str = "teacher",
) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.3)
    for name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"]:
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
        for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            style._element.rPr.rFonts.attrib.pop(qn(f"w:{attribute}"), None)
        for border in style._element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
        style.font.color.rgb = RGBColor(0, 0, 0)
    document.styles["Title"].font.size = Pt(22)
    for name, size in [("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 11), ("Heading 4", 11)]:
        document.styles[name].font.size = Pt(size)
    normal = document.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)
    body = content.body
    teacher = audience == "teacher"
    document.add_heading(content.title, 0)
    document.add_paragraph(f"{'教师版' if teacher else '学生版'}课堂讲义 · {body['duration']} 分钟")
    for key, label in [
        ("objectives", "教学目标"),
        ("key_questions", "重点问题"),
        ("preparation", "课前准备"),
    ]:
        if body.get(key) and (teacher or key != "preparation"):
            document.add_heading(label, 1)
            append_prose(document, body[key])
    document.add_heading("课堂安排", 1)
    elapsed = 0
    for part in body["segments"]:
        document.add_paragraph(f"{elapsed}-{elapsed + part['minutes']} 分钟　{part['title']}")
        elapsed += part["minutes"]
    for index, part in enumerate(body["segments"], 1):
        document.add_heading(f"{index}. {part['title']}", 1)
        document.add_paragraph(
            f"{LESSON_STAGES.get(part.get('stage'), '课堂环节')} · {part['minutes']} 分钟"
        )
        append_prose(document, part["content"])
        for key, label in [
            ("activity", "教师引导与课堂活动"),
            ("example", "例题与演示"),
            ("student_activity", "学生任务"),
            ("expected_response", "预期回答与参考答案"),
            ("check_for_understanding", "检查与反馈"),
            ("misconceptions", "易错点与纠正"),
        ]:
            if part.get(key) and (teacher or key == "student_activity"):
                document.add_heading(label, 2)
                append_prose(document, part[key])
        for quiz_id in part["quiz_ids"]:
            example = body.get("examples", {}).get(quiz_id)
            if example and (teacher or example.get("student_content")):
                document.add_heading(f"关联题组：{example['title']} · v{example['version']}", 2)
                append_prose(document, example["answer"] if teacher else example["student_content"])
        document.add_paragraph(
            "知识点：" + "、".join(node_names[node_id] for node_id in part["node_ids"])
        )
        document.add_paragraph(
            "依据："
            + "、".join(
                f"[{next(i for i, c in enumerate(citations, 1) if c['id'] == chunk_id)}]"
                for chunk_id in part["citation_ids"]
            )
        )
    if teacher and body.get("student_homework"):
        document.add_heading("学生版课后任务", 1)
        append_prose(document, body["student_homework"])
    homework = body.get("homework" if teacher else "student_homework", "")
    if homework:
        document.add_heading("课后巩固与参考要点" if teacher else "课后任务", 1)
        append_prose(document, homework)
    if teacher and body["missing_info"]:
        document.add_heading("待补充事项", 1)
        for item in body["missing_info"]:
            document.add_paragraph(item, style="List Bullet")
    document.add_heading("课程依据", 1)
    for index, source in enumerate(citations, 1):
        document.add_heading(
            f"[{index}] {source['material_title']} · v{source['material_version']} "
            f"· {source['locator']}",
            2,
        )
        for piece in re.split(r"(!\[[^\]]*\]\(asset:rId\d+\))", source["text"]):
            image = re.fullmatch(r"!\[[^\]]*\]\(asset:(rId\d+)\)", piece)
            if image:
                original = Document(material_paths[source["material_id"]])
                image_part = original.part.related_parts[image.group(1)]
                if image_part.content_type not in {"image/png", "image/jpeg"}:
                    document.add_paragraph("此原文插图请下载源文件查看。")
                else:
                    shape = document.add_picture(BytesIO(image_part.blob))
                    scale = min(1, Cm(14) / shape.width, Cm(18) / shape.height)
                    shape.width, shape.height = int(shape.width * scale), int(shape.height * scale)
            else:
                append_prose(document, piece)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
